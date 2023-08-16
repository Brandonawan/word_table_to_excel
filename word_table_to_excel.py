import os
import argparse
from docx import Document
from openpyxl import Workbook

def extract_table_data(docx_path, output_excel_path):
    try:
        doc = Document(docx_path)
        table = doc.tables[0]  # Extract data from the first table

        workbook = Workbook()
        sheet = workbook.active

        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            sheet.append(row_data)

        workbook.save(output_excel_path)
        print(f'Table data extracted and saved to {output_excel_path}')

    except Exception as e:
        print(f'An error occurred: {e}')

def main():
    parser = argparse.ArgumentParser(description='Convert Word table to Excel')
    parser.add_argument('input', help='Path to the input Word document')
    parser.add_argument('-o', '--output', default='table_data.xlsx', help='Path to the output Excel file')
    args = parser.parse_args()

    input_path = args.input
    output_path = args.output

    if not os.path.isfile(input_path):
        print('Error: Input file not found')
        return

    extract_table_data(input_path, output_path)

if __name__ == '__main__':
    main()
